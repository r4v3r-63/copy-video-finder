import json
import re

import pytest

from app import catalog as cat
from app import licenses, paths, processes
from app.export import html_catalog, registry, summary
from app.sources.base import LicenseInfo, VideoFile, VideoResult
from tests.fakes import Clock


def video(sid, title="Ламинатор в офисе", source="Pixabay", lic="Pixabay Content License", status=licenses.ALLOWED):
    return VideoResult(source=source, source_id=sid, title=title, description="desc; с точкой с запятой", preview_url="https://p/1.jpg",
                       duration=12.4, width=1920, height=1080, has_audio=True, page_url=f"https://page/{sid}",
                       files=[VideoFile("Full HD", f"https://cdn/{sid}.mp4", 1920, 1080, "video/mp4")],
                       author="Анна", author_url="https://a",
                       license=LicenseInfo(lic, "https://lic", status=status, note=""))


@pytest.fixture
def clock():
    return Clock()


@pytest.fixture
def cats():
    return processes.load_catalog(paths.resources_dir() / "processes.json")


@pytest.fixture
def store(tmp_path, clock):
    s = cat.Catalog(tmp_path / "c.sqlite", now=clock)
    a = s.select(video("1"), "lam_pouch", cat.EXACT)
    s.link(a, "lam_roll", cat.RELATED)
    s.set_comment(a, "в холл")
    s.select(video("2", title="Брошюровщик", source="Wikimedia Commons", lic="CC BY-SA 4.0",
                   status=licenses.ATTRIBUTION), "comb_binding", cat.PRODUCT)
    s.add(video("3", title="Без услуги"))
    s.reject(video("4", title="Пляж"))
    yield s
    s.close()


def procs(cats):
    return processes.all_processes(cats)


def test_registry_rows_one_per_link(store, cats):
    rows = registry.rows(store, procs(cats))
    keyed = [(r["id"], r["service_id"], r["match"], r["counted"]) for r in rows]
    assert keyed == [(1, "lam_pouch", cat.EXACT, "да"), (1, "lam_roll", cat.RELATED, "да"),
                     (2, "comb_binding", cat.PRODUCT, "нет"), (3, "", "", "нет"), (4, "", "", "нет")]
    first = rows[0]
    assert first["category"] == "Ламинация" and first["service"] == "Пакетная ламинация"
    assert first["license_confirmed"] == "да" and first["comment"] == "в холл" and first["has_audio"] == "да"
    assert rows[2]["attribution"] == "Анна — Брошюровщик — CC BY-SA 4.0 — https://page/2"


def test_csv_for_russian_excel_and_round_trip(store, cats, tmp_path):
    path = tmp_path / "реестр.csv"
    registry.write_csv(path, registry.rows(store, procs(cats)))
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM — Excel узнаёт UTF-8
    assert raw.decode("utf-8-sig").splitlines()[0].startswith("ID в каталоге;Категория;Услуга")
    # Перенос на «другой компьютер» — в пустой каталог.
    other = cat.Catalog(tmp_path / "other.sqlite")
    videos, links = registry.import_csv(path, other, procs(cats))
    assert (videos, links) == (4, 3)
    e1 = other.find(("Pixabay", "1"))
    assert e1.links == {"lam_pouch": cat.EXACT, "lam_roll": cat.RELATED}
    assert e1.status == cat.SELECTED and e1.license_confirmed and e1.comment == "в холл"
    assert e1.result.files[0].url == "https://cdn/1.mp4" and e1.result.duration == 12.4
    assert e1.result.description == "desc; с точкой с запятой"
    assert other.find(("Pixabay", "4")).status == cat.REJECTED
    assert other.coverage() == store.coverage()
    # Повторный импорт не создаёт дублей.
    registry.import_csv(path, other, procs(cats))
    assert len(other.entries()) == 4
    other.close()


def test_import_recomputes_license_status(tmp_path, cats):
    path = tmp_path / "r.csv"
    path.write_text("ID в каталоге;Площадка;ID на площадке;Лицензия;Статус лицензии;Код услуги\n"
                    "1;Vimeo;9;CC BY-NC 4.0;Разрешено;lam_pouch\n", encoding="utf-8-sig")
    store = cat.Catalog(tmp_path / "c.sqlite")
    registry.import_csv(path, store, procs(cats))
    assert store.find(("Vimeo", "9")).result.license.status == licenses.REJECTED  # из файла не берётся на веру
    store.close()


def test_xlsx(store, cats, tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "реестр.xlsx"
    s = summary.build(store, cats)
    registry.write_xlsx(path, registry.rows(store, procs(cats)), summary.lines(s))
    wb = load_workbook(path)
    ws = wb["Реестр"]
    assert ws.cell(1, 1).value == "ID в каталоге" and ws.max_row == 6
    page_col = registry.HEADERS.index("Страница ролика") + 1
    assert ws.cell(2, page_col).hyperlink.target == "https://page/1"
    assert wb["Сводка покрытия"].cell(1, 1).value == "Сводка покрытия"


def test_summary_and_added_since_previous(store, cats, clock, tmp_path):
    s = summary.build(store, cats, now=clock)
    assert s["stats"]["counted"] == 2 and len(s["added"]) == 3 and s["since"] is None
    need = [p.id for p, _ in s["need_search"]]
    assert "lam_pouch" in need and need.index("lam_wide") < need.index("lam_pouch")  # без роликов — раньше
    summary.write_html(tmp_path / "сводка.html", s)
    assert "Нужен дополнительный поиск" in (tmp_path / "сводка.html").read_text(encoding="utf-8")
    summary.mark_done(store, now=clock)
    clock.t += 100
    store.select(video("5", title="Новый"), "lam_wide")
    s2 = summary.build(store, cats, now=clock)
    assert [e.result.source_id for e in s2["added"]] == ["5"] and s2["since"]


def test_html_catalog_is_self_contained(store, cats, tmp_path):
    store.select(video("6", title="Опасное </script><b>название</b>"), "lam_pouch")
    path = tmp_path / "каталог.html"
    count = html_catalog.write(path, store, cats)
    html = path.read_text(encoding="utf-8")
    assert count == 4  # отклонённые не выгружаются
    assert "<script src" not in html and "<link" not in html  # без внешних файлов
    data = json.loads(re.search(r"const DATA = (.*?);\n", html).group(1).replace("<\\/", "</"))
    assert {i["id"] for i in data["items"]} == {1, 2, 3, 5}
    assert html.count("</script>") == 1  # опасное название не закрыло тег
    item = next(i for i in data["items"] if i["id"] == 1)
    assert item["services"][0] == {"category": "Ламинация", "name": "Пакетная ламинация", "match": cat.EXACT}
    assert item["files"][0]["url"] == "https://cdn/1.mp4"
