"""Реестр каталога: CSV и Excel, строка на каждую привязку «ролик — услуга», и импорт CSV.

CSV — UTF-8 с меткой BOM и разделителем «;»: так его без настройки открывает
русский Excel. Тот же CSV можно загрузить обратно (перенос каталога на другой компьютер).
"""

import csv
import json
import math
from datetime import datetime
from pathlib import Path

from app import paths
from app.core import licenses
from app.core.models import LicenseInfo, VideoFile, VideoResult
from app.storage import catalog as cat

# (ключ, заголовок столбца)
COLUMNS = [
    ("id", "ID в каталоге"),
    ("category", "Категория"),
    ("service", "Услуга"),
    ("service_id", "Код услуги"),
    ("match", "Степень соответствия"),
    ("counted", "Засчитан в покрытие"),
    ("source", "Площадка"),
    ("source_id", "ID на площадке"),
    ("title", "Название"),
    ("description", "Описание"),
    ("duration", "Длительность, с"),
    ("width", "Ширина"),
    ("height", "Высота"),
    ("has_audio", "Звук"),
    ("license", "Лицензия"),
    ("license_status", "Статус лицензии"),
    ("license_confirmed", "Лицензия подтверждена"),
    ("license_url", "Условия лицензии"),
    ("license_note", "Пояснение к лицензии"),
    ("author", "Автор"),
    ("author_url", "Ссылка на автора"),
    ("attribution", "Текст атрибуции"),
    ("page_url", "Страница ролика"),
    ("preview_url", "Превью"),
    ("files", "Файлы (из API)"),
    ("status", "Отметка"),
    ("downloaded_path", "Скачан в"),
    ("comment", "Комментарий"),
    ("added_at", "Добавлен"),
    ("checked_at", "Лицензия проверена"),
]
HEADERS = [title for _, title in COLUMNS]
_BY_TITLE = {title: key for key, title in COLUMNS}
YES, NO = "да", "нет"


def _date(ts: float | None) -> str:
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M") if ts else ""


def _yes_no(value) -> str:
    return "" if value is None else (YES if value else NO)


def rows(store: cat.Catalog, processes) -> list[dict]:
    """Строки реестра: по одной на каждую привязку; ролик без услуги — одной строкой."""
    by_id = {p.id: p for p in processes}
    result = []
    for e in sorted(store.entries(), key=lambda e: e.id):
        r, lic = e.result, e.result.license
        base = {
            "id": e.id, "source": r.source, "source_id": r.source_id, "title": r.title,
            "description": r.description, "duration": "" if r.duration is None else round(r.duration, 1),
            "width": r.width or "", "height": r.height or "", "has_audio": _yes_no(r.has_audio),
            "license": lic.name, "license_status": lic.status, "license_confirmed": _yes_no(e.license_confirmed),
            "license_url": lic.url or "", "license_note": lic.note, "author": r.author,
            "author_url": r.author_url or "",
            "attribution": licenses.attribution_text(r.author, r.title, lic.name, r.page_url)
            if lic.status == licenses.ATTRIBUTION else "",
            "page_url": r.page_url, "preview_url": r.preview_url or "",
            "files": json.dumps([f.__dict__ for f in r.files], ensure_ascii=False) if r.files else "",
            "status": e.status, "downloaded_path": paths.portable_path(e.downloaded_path), "comment": e.comment,
            "added_at": _date(e.added_at), "checked_at": _date(e.checked_at),
        }
        links = sorted(e.links.items()) or [(None, None)]
        for process_id, match in links:
            proc = by_id.get(process_id)
            counted = e.counts_for_coverage and match not in (None, cat.PRODUCT)
            result.append({**base, "category": proc.category if proc else "", "service": proc.name if proc else "",
                           "service_id": process_id or "", "match": match or "", "counted": _yes_no(counted)})
    return result


# Строка с таким началом в Excel — формула. Названия и описания приходят с площадок, поэтому
# в CSV перед ними ставится апостроф (при импорте он снимается), а в XLSX они пишутся как текст.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
REQUIRED = ("source", "source_id")


def _csv_safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(_FORMULA_START) else value


def _csv_value(value: str) -> str:
    return value[1:] if value.startswith("'") and value[1:].startswith(_FORMULA_START) else value


def write_csv(path: Path, data: list[dict]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(HEADERS)
        for row in data:
            writer.writerow([_csv_safe(row.get(key, "")) for key, _ in COLUMNS])


def write_xlsx(path: Path, data: list[dict], summary: list[tuple[str, str]] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in data:
        ws.append([row.get(key, "") for key, _ in COLUMNS])
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":  # строка с «=» — текст с площадки, а не формула
                cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"title": 40, "description": 40, "service": 32, "license_note": 30, "attribution": 40, "files": 30,
              "comment": 30, "page_url": 30}
    for i, (key, title) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, max(10, len(title) + 2))
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
    if summary:
        ws2 = wb.create_sheet("Сводка покрытия")
        for line in summary:
            ws2.append(list(line))
        ws2.column_dimensions["A"].width = 50
        ws2.column_dimensions["B"].width = 60
    wb.save(path)


def import_csv(path: Path, store: cat.Catalog, processes) -> tuple[int, int]:
    """Загрузить реестр CSV: (роликов добавлено или обновлено, привязок к услугам).

    Отметки, подтверждение лицензии и комментарий берутся из файла; степень соответствия —
    тоже. Статус лицензии пересчитывается по таблице лицензий, а не берётся на веру.
    """
    known = {p.id for p in processes}
    videos, links = set(), 0
    with open(path, encoding="utf-8-sig", newline="") as f:
        delimiter = _delimiter(f.readline())
        f.seek(0)
        reader = csv.DictReader(f, delimiter=delimiter)
        columns = {_BY_TITLE.get(k.strip(), k.strip()) for k in reader.fieldnames or [] if k}
        missing = [title for key, title in COLUMNS if key in REQUIRED and key not in columns]
        if missing:
            raise ValueError("это не реестр программы — нет столбцов: " + ", ".join(missing))
        for raw in reader:
            row = {_BY_TITLE.get(k.strip(), k.strip()): _csv_value((v or "").strip())
                   for k, v in raw.items() if k and isinstance(v, str)}
            if not row.get("source") or not row.get("source_id"):
                continue
            try:
                files = [VideoFile(**f) for f in json.loads(row.get("files") or "[]")]
            except (ValueError, TypeError):
                files = []
            status, note = licenses.evaluate(row.get("license"))
            result = VideoResult(
                source=row["source"], source_id=row["source_id"], title=row.get("title", ""),
                description=row.get("description", ""), preview_url=row.get("preview_url") or None,
                duration=_float(row.get("duration")), width=_int(row.get("width")), height=_int(row.get("height")),
                has_audio={YES: True, NO: False}.get(row.get("has_audio", "")), page_url=row.get("page_url", ""),
                files=files, author=row.get("author", ""), author_url=row.get("author_url") or None,
                license=LicenseInfo(name=row.get("license", ""), url=row.get("license_url") or None, status=status,
                                    note=row.get("license_note") or note),
            )
            video_id = store.add(result)
            videos.add(video_id)
            if row.get("status") in (cat.SELECTED, cat.REJECTED, cat.NEW):
                store.set_status(video_id, row["status"])
            # Столбец есть — значение из файла, даже «нет» и пустое; столбца нет — в каталоге как было.
            if "license_confirmed" in columns:
                store.confirm_license(video_id, row.get("license_confirmed") == YES)
            if "comment" in columns:
                store.set_comment(video_id, row.get("comment", ""))
            if "downloaded_path" in columns:
                store.set_downloaded(video_id, row.get("downloaded_path", ""))
            service = row.get("service_id", "")
            if service in known:
                match = row.get("match") if row.get("match") in cat.MATCHES else cat.EXACT
                store.link(video_id, service, match)
                links += 1
    return len(videos), links


def _delimiter(header: str) -> str:
    """Разделитель — по строке заголовков: в ней нет данных с запятыми, как в названиях и тегах."""
    for delimiter in (";", ",", "\t"):
        names = next(csv.reader([header], delimiter=delimiter), [])
        if sum(name.strip() in _BY_TITLE for name in names) >= 2:
            return delimiter
    return ";" if header.count(";") >= header.count(",") else ","


def _float(text) -> float | None:
    try:
        value = float(str(text).replace(",", ".")) if str(text).strip() else None
        return value if value is not None and math.isfinite(value) else None
    except ValueError:
        return None


def _int(text) -> int | None:
    value = _float(text)
    return int(value) if value is not None else None
